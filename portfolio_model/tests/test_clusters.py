"""Correlation clusters: readable labels and the cluster table surfaced to readers."""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

PKG_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PKG_ROOT))

from src.risk.risk import cluster_summary, correlation_clusters  # noqa: E402


def _corr(pairs: dict[tuple[str, str], float], tickers: list[str]) -> pd.DataFrame:
    df = pd.DataFrame(np.eye(len(tickers)), index=tickers, columns=tickers)
    for (a, b), v in pairs.items():
        df.loc[a, b] = df.loc[b, a] = v
    return df


CORR = _corr({("ADBE", "DIS"): 0.86, ("DIS", "DUOL"): 0.78, ("VOO", "ADBE"): 0.42},
             ["ADBE", "DIS", "DUOL", "VOO", "CASH"])
WEIGHTS = pd.Series({"ADBE": 0.20, "DIS": 0.15, "DUOL": 0.10, "VOO": 0.40, "CASH": 0.15})


def test_clusters_are_reported_with_readable_labels():
    clusters = correlation_clusters(CORR, 0.70)
    # single linkage joins ADBE-DIS and DIS-DUOL into one named group
    assert clusters["ADBE"] == clusters["DIS"] == clusters["DUOL"]
    assert clusters["ADBE"].startswith("cluster:")
    assert set(clusters["ADBE"].removeprefix("cluster:").split(",")) == {"ADBE", "DIS", "DUOL"}
    assert clusters["VOO"] == "independent"
    assert clusters["CASH"] == "independent"


def test_cluster_table_shows_members_and_combined_weight():
    table = cluster_summary(CORR, WEIGHTS, 0.70)
    assert len(table) == 1
    row = table.iloc[0]
    assert row["members"] == "ADBE, DIS, DUOL"
    assert row["n_members"] == 3
    assert row["combined_weight"] == pytest.approx(0.45)
    assert 0.70 <= row["mean_pairwise_correlation"] <= 1.0 or row["n_members"] > 2
    # single linkage can join two strongly correlated pairs whose cross-pair is weak,
    # so the mean across all members may sit below the threshold. That is the point of
    # showing the mean: it tells the reader how tight the cluster really is.
    assert row["mean_pairwise_correlation"] == pytest.approx((0.86 + 0.78 + 0.0) / 3)
    assert "one exposure" in row["note"]


def test_cluster_table_is_empty_when_nothing_is_correlated():
    table = cluster_summary(_corr({("ADBE", "DIS"): 0.20}, ["ADBE", "DIS"]),
                           pd.Series({"ADBE": 0.5, "DIS": 0.5}), 0.70)
    assert len(table) == 0


def test_pipeline_surfaces_clusters_in_the_workbook_and_payload():
    import tempfile

    from openpyxl import load_workbook

    from src.pipeline import run
    from web.serialize import build_payload

    out = Path(tempfile.mkdtemp())
    res = run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_full.xlsx"),
              out_dir=str(out),
              overrides={"sims": 400, "sensitivity_sims": 200, "offline": "always"},
              render_charts=False)
    assert "cluster_sheet" in res
    payload = build_payload(res)
    assert "clusters" in payload["tables"]
    ws = load_workbook(out / "spreadsheet" / "portfolio_analysis.xlsx")["Risk"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Correlated Clusters" in text, "the Risk sheet should show the cluster table"
