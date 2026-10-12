"""Tests for the interpretability layer: glossary completeness, generated docs, workbook sheet.

The point of these tests is that the *explanations* cannot silently drift away from the code.
"""
from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

PKG_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PKG_ROOT))

from src.metric_glossary import GLOSSARY, SHEET_GUIDE, SHEET_QUESTIONS  # noqa: E402

FIELDS = ("label", "unit", "formula", "meaning", "read_as", "action", "caveat")
SUBSTANTIVE = ("formula", "meaning", "read_as", "action", "caveat")


def _load_generator():
    spec = importlib.util.spec_from_file_location(
        "build_definitions_doc", PKG_ROOT / "scripts" / "build_definitions_doc.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_every_glossary_entry_is_complete():
    assert len(GLOSSARY) >= 40, "the glossary should cover the whole output surface"
    for key, entry in GLOSSARY.items():
        for f in FIELDS:
            value = entry.get(f)
            assert isinstance(value, str) and value.strip(), f"{key} is missing '{f}'"
            if f in SUBSTANTIVE:
                assert len(value) > 15, f"{key}.{f} is too thin to be useful"
        assert entry["unit"] in {"pct", "money", "ratio", "count", "years", "days", "text", "score"}, \
            f"{key} has an unsupported unit"


def test_definitions_markdown_is_generated_and_up_to_date(tmp_path):
    gen = _load_generator()
    doc = gen.build_doc()
    for key, entry in GLOSSARY.items():
        assert f"`{key}`" in doc, f"{key} is missing from the generated definitions page"
        assert entry["label"] in doc
    assert "misreading" in doc.lower()

    out = tmp_path / "METRIC_DEFINITIONS.md"
    out.write_text(doc, encoding="utf-8")
    assert out.read_text(encoding="utf-8") == gen.build_doc()


def test_committed_definitions_page_matches_the_glossary():
    """docs/METRIC_DEFINITIONS.md must not drift from src/metric_glossary.py."""
    committed = PKG_ROOT.parent / "docs" / "METRIC_DEFINITIONS.md"
    assert committed.exists(), "docs/METRIC_DEFINITIONS.md is missing"
    text = committed.read_text(encoding="utf-8")
    missing = [k for k in GLOSSARY if f"`{k}`" not in text]
    assert not missing, f"definitions page is out of date, missing: {missing}"


def test_workbook_contains_definitions_sheet_with_every_metric(tmp_path):
    from src.reporting.excel import build_workbook

    results = {
        "meta": {"generated_utc": "x", "seed": 1, "n_sims": 10, "version": "t",
                 "portfolio_name": "t", "input": "t", "python": "3", "platform": "t",
                 "weights_source": "supplied", "data_as_of": "2026-01-01",
                 "data_sources": "bundled snapshot", "offline_mode": "auto"},
        "dashboard": [("Metric", 1.0, "0.0")],
        "holdings_sheet": __import__("pandas").DataFrame({"ticker": ["A"], "weight": [1.0]}),
        "sector_sheet": __import__("pandas").DataFrame({"weight": [1.0], "n_holdings": [1]}),
        "corr": __import__("pandas").DataFrame({"A": [1.0]}, index=["A"]),
        "history_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "mc": {"model": "gaussian", "n_sims": 10}, "mc_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "mc_percentiles": __import__("pandas").DataFrame({"a": [1]}),
        "markov_states": __import__("pandas").DataFrame({"a": [1]}),
        "markov_P": __import__("pandas").DataFrame({"a": [1]}),
        "markov_misc": __import__("pandas").DataFrame({"a": [1]}),
        "bs": {"method": "block", "block_size_days": 20, "n_sims": 10},
        "bs_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "comparison_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "disagreement_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "sensitivity_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "risk_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "recovery_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "stress_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "hypothetical_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "glide_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "glide_impact_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "rebalance_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "rule_flags": [], "thesis_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "dq_sheet": __import__("pandas").DataFrame({"a": [1]}),
        "scorecard": __import__("pandas").DataFrame({"a": [1]}),
        "decision_matrix": __import__("pandas").DataFrame({"a": [1]}),
        "stop_sheet": __import__("pandas").DataFrame({"a": [1]}),
    }
    path = build_workbook(tmp_path / "defs_check.xlsx", results)
    from openpyxl import load_workbook

    ws = load_workbook(path)["Definitions"]
    keys = {row[0].value for row in ws.iter_rows(min_row=4)}
    assert set(GLOSSARY) <= keys, f"workbook Definitions sheet is missing {set(GLOSSARY) - keys}"


def test_sensitivity_fragility_threshold_is_read_at_the_base_return_assumption():
    import pandas as pd

    from src.sensitivity.sensitivity import interpret_grid

    rows = []
    for vm, succ in ((0.7, 0.99), (0.8, 0.97), (0.9, 0.95), (1.0, 0.92), (1.1, 0.90),
                     (1.2, 0.85), (1.3, 0.74), (1.4, 0.70), (1.5, 0.65)):
        for rs in (-0.04, -0.02, 0.0, 0.02, 0.04):
            # pessimistic return shifts fail early; the fragility read must ignore them
            s = succ if rs == 0.0 else max(0.2, succ - 0.25)
            rows.append({"return_shift": rs, "vol_multiplier": vm, "prob_success": s,
                         "median_terminal": 1000.0})
    out = interpret_grid(pd.DataFrame(rows))
    assert out["fragile_vol_threshold"] == 1.3, out


def test_report_cross_references_point_at_sections_that_actually_exist():
    """The 'How to read this report' table must not point at section numbers that are not there.

    Renumbering a report is a silent failure: the prose keeps citing the old numbers and the
    reader follows a reference to a section about something else entirely.
    """
    import re
    import tempfile

    from docx import Document

    from src.pipeline import run

    out = Path(tempfile.mkdtemp())
    run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_minimal.xlsx"),
        out_dir=str(out), overrides={"sims": 300, "sensitivity_sims": 200, "offline": "always"},
        render_charts=False)
    doc = Document(out / "report" / "portfolio_analysis_report.docx")

    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    numbered = {}
    for h in headings:
        m = re.match(r"^(\d+(?:\.\d+)?)\.?\s+(.*)$", h)
        if m:
            numbered[m.group(1)] = m.group(2).lower()

    guide = next(t for t in doc.tables
                 if t.rows and t.rows[0].cells[0].text == "question")
    refs = [(r.cells[0].text, r.cells[1].text) for r in guide.rows][1:]
    assert len(refs) >= 6, "the reading guide should cover the whole report"

    for question, target in refs:
        for num in re.findall(r"Section (\d+(?:\.\d+)?)", target):
            assert num in numbered, f"{question!r} points at Section {num}, which does not exist"
            named = target.split(f"Section {num}", 1)[1].split(";")[0].strip().lower()
            assert named.split(" ")[0] in numbered[num], (
                f"{question!r} cites 'Section {num} {named}' but that section is '{numbered[num]}'")
        if "Appendix A" in target:
            assert any("Appendix A" in h for h in headings), "Appendix A is referenced but not written"


def test_workbook_opens_on_a_read_me_sheet_that_documents_every_sheet():
    """A 21-sheet workbook is useless if the reader cannot find the right sheet."""
    import tempfile

    from openpyxl import load_workbook

    from src.pipeline import run

    out = Path(tempfile.mkdtemp())
    run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_minimal.xlsx"),
        out_dir=str(out), overrides={"sims": 300, "sensitivity_sims": 200, "offline": "always"},
        render_charts=False)
    wb = load_workbook(out / "spreadsheet" / "portfolio_analysis.xlsx")

    assert wb.sheetnames[0] == "Read Me", f"workbook should open on Read Me, got {wb.sheetnames[0]}"
    documented = dict(SHEET_GUIDE)
    undocumented = [s for s in wb.sheetnames if s != "Read Me" and s not in documented]
    assert not undocumented, f"sheets with no entry in SHEET_GUIDE: {undocumented}"
    stale = [s for s in documented if s not in wb.sheetnames]
    assert not stale, f"SHEET_GUIDE documents sheets that no longer exist: {stale}"

    text = " ".join(str(c.value) for row in wb["Read Me"].iter_rows(max_col=2) for c in row if c.value)
    for question, _ in SHEET_QUESTIONS:
        assert question in text, f"Read Me sheet is missing the question {question!r}"


def test_severity_legend_covers_every_severity_the_rule_engine_can_emit():
    """If the rule engine gains a severity, the legend has to explain it."""
    import re

    from src.metric_glossary import SEVERITY_LEGEND

    legend = {s for s, _ in SEVERITY_LEGEND}
    assert {"CRITICAL", "HIGH", "WARNING", "MONITOR", "INFO"} <= legend

    source = (PKG_ROOT / "src" / "rules" / "rule_engine.py").read_text(encoding="utf-8")
    emitted = set(re.findall(r"severity\s*=\s*[\"']([A-Z]+)[\"']", source))
    emitted |= set(re.findall(r"[\"'](CRITICAL|HIGH|WARNING|MONITOR|INFO)[\"']", source))
    unexplained = emitted - legend
    assert not unexplained, f"rule engine emits {sorted(unexplained)} but no legend explains them"


def test_report_contains_every_numbered_section_it_claims():
    """The READMEs promise 23 numbered sections; a section silently dropped is a hole in the report."""
    import re
    import tempfile

    from docx import Document

    from src.pipeline import run

    out = Path(tempfile.mkdtemp())
    run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_minimal.xlsx"),
        out_dir=str(out), overrides={"sims": 300, "sensitivity_sims": 200, "offline": "always"},
        render_charts=False)
    doc = Document(out / "report" / "portfolio_analysis_report.docx")
    top = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    numbers = {int(m.group(1)) for m in (re.match(r"^(\d+)(?:\.\d+)?\.?\s", h) for h in top) if m}
    missing = sorted(set(range(1, 24)) - numbers)
    assert not missing, f"report is missing numbered sections {missing}"
    for must in ("How to read this report", "Final Conclusions", "Monitoring Checklist",
                 "Appendix A. Metric Definitions"):
        assert must in top, f"report is missing the {must!r} section"


def test_every_rule_flag_points_at_a_real_metric_definition():
    """A suggestion is only useful if the metric behind it is defined somewhere."""
    from src.metric_glossary import GLOSSARY
    from src.rules.rule_engine import FLAG_GLOSSARY, glossary_key_for

    dangling = {prefix: key for prefix, key in FLAG_GLOSSARY.items() if key not in GLOSSARY}
    assert not dangling, f"flag families point at undefined metrics: {dangling}"
    for fid in ("POS-01", "GRP-02", "RESV-05", "MODEL-06", "THESIS-01", "TAIL-01", "SENS-01"):
        assert glossary_key_for(fid) in GLOSSARY, f"{fid} has no glossary entry"


def test_report_explains_each_high_severity_flag_in_place():
    """Section 19.6 must define the metric, not just restate the number."""
    import tempfile

    from docx import Document

    from src.pipeline import run

    out = Path(tempfile.mkdtemp())
    res = run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_full.xlsx"),
              out_dir=str(out), overrides={"sims": 400, "sensitivity_sims": 200, "offline": "always"},
              render_charts=False)
    doc = Document(out / "report" / "portfolio_analysis_report.docx")
    paras = [p.text for p in doc.paragraphs]
    assert any(t.startswith("19.6") for t in paras), "report has no per-flag explanation section"

    high = [f for f in res["rule_flags"] if f["severity"] in ("CRITICAL", "HIGH")]
    assert high, "expected at least one high-severity flag for the example portfolio"
    for f in high:
        assert f["id"] in " ".join(paras), f"flag {f['id']} is not explained in the report"
    joined = " ".join(paras)
    assert joined.count("Metric: ") >= len(high), "flags are listed without defining their metric"
    assert joined.count("The trap: ") >= len(high), "flags are listed without their misreading warning"


def test_workbook_rule_flags_sheet_carries_the_metric_definition():
    import tempfile

    from openpyxl import load_workbook

    from src.pipeline import run

    out = Path(tempfile.mkdtemp())
    run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_full.xlsx"),
        out_dir=str(out), overrides={"sims": 400, "sensitivity_sims": 200, "offline": "always"},
        render_charts=False)
    ws = load_workbook(out / "spreadsheet" / "portfolio_analysis.xlsx")["Rule Flags"]
    header_row = next(r for r in ws.iter_rows(max_col=12) if r[0].value == "id")
    headers = [c.value for c in header_row if c.value]
    assert "metric_definition" in headers, f"Rule Flags sheet has no definition column: {headers}"
    first_data = next(ws.iter_rows(min_row=header_row[0].row + 1, max_col=len(headers)))
    values = dict(zip(headers, [c.value for c in first_data]))
    assert values.get("metric_definition"), "the definition column is empty on the first flag"


WORKED_FIELDS = ("key", "metric", "unit", "formula", "your_numbers", "result",
                 "what_it_means", "the_trap", "what_to_do")


def _worked_examples_for_example_portfolio():
    import tempfile

    from src.interpret import worked_examples
    from src.pipeline import run

    res = run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_full.xlsx"),
              out_dir=str(Path(tempfile.mkdtemp())),
              overrides={"sims": 500, "sensitivity_sims": 200, "offline": "always"},
              render_charts=False)
    return res, worked_examples(res)


def test_worked_examples_cover_the_metrics_that_drive_the_decision():
    _, examples = _worked_examples_for_example_portfolio()
    keys = {e["key"] for e in examples}
    required = {"weights_source", "hhi", "effective_n", "top3_top5_top10", "correlation_cluster",
                "required_reserve", "terminal_percentiles", "prob_terminal_gt_start",
                "prob_all_payments_funded", "cvar", "max_drawdown_growth", "glide_path",
                "rebalance_band", "engine_disagreement", "data_provenance"}
    missing = required - keys
    assert not missing, f"no worked example produced for: {missing}"


def test_every_worked_example_is_complete_and_uses_a_defined_metric():
    _, examples = _worked_examples_for_example_portfolio()
    assert examples, "no worked examples produced"
    for e in examples:
        assert e["key"] in GLOSSARY, f"worked example for undefined metric {e['key']}"
        for field in WORKED_FIELDS:
            assert str(e.get(field, "")).strip(), f"{e['key']} is missing {field}"
        assert len(e["your_numbers"]) > 15, f"{e['key']} shows no arithmetic"


def test_worked_examples_quote_this_portfolios_actual_holdings():
    res, examples = _worked_examples_for_example_portfolio()
    tickers = list(res["holdings"]["ticker"])
    text = " ".join(e["your_numbers"] for e in examples)
    quoted = [t for t in tickers if t in text]
    assert len(quoted) >= 3, f"worked examples do not quote the real holdings: {text[:200]}"
    # the HHI line must be the sum of the actual squared weights
    hhi = next(e for e in examples if e["key"] == "hhi")
    for _, r in res["holdings"].iterrows():
        assert f"{float(r['weight']):.3f}²" in hhi["your_numbers"], (
            f"HHI worked example omits {r['ticker']} at {r['weight']:.3f}: {hhi['your_numbers']}")


def test_worked_examples_reach_the_workbook_and_the_report():
    import tempfile

    from openpyxl import load_workbook

    from src.pipeline import run

    out = Path(tempfile.mkdtemp())
    run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_full.xlsx"),
        out_dir=str(out), overrides={"sims": 400, "sensitivity_sims": 200, "offline": "always"},
        render_charts=False)

    ws = load_workbook(out / "spreadsheet" / "portfolio_analysis.xlsx")["Worked Examples"]
    body = " ".join(str(c.value) for row in ws.iter_rows(max_col=8) for c in row if c.value)
    assert "Calculation" in body or "your_numbers" in body, "Worked Examples sheet is empty"
    assert "²" in body, "the HHI arithmetic is not shown in the workbook"

    from docx import Document
    doc = Document(out / "report" / "portfolio_analysis_report.docx")
    paras = [p.text for p in doc.paragraphs]
    assert any(t.startswith("Appendix B") for t in paras), "report has no worked-examples appendix"
    i = next(n for n, t in enumerate(paras) if t.startswith("Appendix B"))
    assert any("Calculation:" in t for t in paras[i:])


def test_metric_values_binds_definitions_to_this_runs_numbers():
    import tempfile

    from src.interpret import metric_values
    from src.pipeline import run

    res = run(input_path=str(PKG_ROOT / "data" / "input" / "example_portfolio_full.xlsx"),
              out_dir=str(Path(tempfile.mkdtemp())),
              overrides={"sims": 400, "sensitivity_sims": 200, "offline": "always"},
              render_charts=False)
    values = metric_values(res)
    assert len(values) >= 20, f"only {len(values)} metrics bound to the run"
    undefined = [k for k in values if k not in GLOSSARY]
    assert not undefined, f"metric values reference undefined glossary keys: {undefined}"

    assert values["required_reserve"].startswith("$")
    assert values["prob_all_payments_funded"].count("%") == 2, "both engines should be shown"
    assert "·" in values["terminal_percentiles"]
    # the bound value must match the run, not be a generic string
    assert values["hhi"] == f"{res['conc']['hhi']['raw']:.4f}"
    assert values["effective_n"] == f"{res['conc']['hhi']['effective_n']:.2f}"

    payload_values = build_payload_for(res)["metric_values"]
    assert payload_values == values


def build_payload_for(res):
    from web.serialize import build_payload
    return build_payload(res)
