"""Tests for the tutorial layer: example workbooks, docs, and the web UI's interpretability routes."""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

PKG_ROOT = Path(__file__).resolve().parent.parent
REPO_ROOT = PKG_ROOT.parent
DOCS = REPO_ROOT / "docs"
sys.path.insert(0, str(PKG_ROOT))

from src.config import load_config  # noqa: E402
from src.ingestion.ingest import ingest  # noqa: E402
from src.validation.validate import validate  # noqa: E402

EXAMPLES = {
    "example_portfolio_minimal.xlsx": {"CASH", "IEF", "VOO", "XOM"},
    "example_portfolio_full.xlsx": {"ADBE", "CASH", "DIS", "DUOL", "IEF", "SHY", "VOO", "XOM"},
}


@pytest.mark.parametrize("name,expected", sorted(EXAMPLES.items()))
def test_example_workbooks_ingest_and_validate(name, expected):
    path = PKG_ROOT / "data" / "input" / name
    assert path.exists(), f"{name} is missing from data/input"
    res = ingest(str(path))
    assert set(res.holdings["ticker"]) == expected
    val = validate(res.holdings, load_config())
    assert not val.errors, f"{name} failed validation: {val.errors}"
    assert abs(val.holdings["weight"].sum() - 1.0) < 1e-9


@pytest.mark.parametrize("name", sorted(EXAMPLES))
def test_example_workbooks_carry_an_instructions_sheet(name):
    from openpyxl import load_workbook

    wb = load_workbook(PKG_ROOT / "data" / "input" / name)
    assert "Instructions" in wb.sheetnames, f"{name} has no Instructions sheet"
    text = " ".join(str(c.value) for row in wb["Instructions"].iter_rows() for c in row if c.value)
    for must in ("Ticker", "Weight", "python main.py", "METRIC_DEFINITIONS"):
        assert must in text, f"Instructions sheet in {name} does not mention {must}"


def test_example_workbook_builder_is_reproducible(tmp_path):
    """The example workbooks are generated from a spec, so regenerating must not change holdings."""
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "build_example_workbooks", PKG_ROOT / "scripts" / "build_example_workbooks.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.build(tmp_path)
    for name, expected in EXAMPLES.items():
        assert set(ingest(str(tmp_path / name)).holdings["ticker"]) == expected


def test_tutorial_and_definitions_pages_exist():
    tutorial = DOCS / "TUTORIAL.md"
    definitions = DOCS / "METRIC_DEFINITIONS.md"
    assert tutorial.exists() and definitions.exists()
    text = tutorial.read_text(encoding="utf-8")
    for must in ("--check", "--offline", "requirements.txt", "METRIC_DEFINITIONS.md",
                 "example_portfolio_full.xlsx", "Troubleshooting"):
        assert must in text, f"TUTORIAL.md does not cover {must}"
    assert len(text) > 4000, "the tutorial is too thin to be a tutorial"


def test_web_ui_exposes_interpretability_routes():
    from web.app import app

    paths = {r.path for r in app.routes}
    for route in ("/api/glossary", "/api/selfcheck", "/api/guide", "/api/analyze",
                  "/api/results/{job_id}", "/api/export/{job_id}/{kind}"):
        assert route in paths, f"missing route {route}"


def test_web_ui_renders_definitions_and_tutorial_tabs():
    js = (PKG_ROOT / "web" / "static" / "app.js").read_text(encoding="utf-8")
    html = (PKG_ROOT / "web" / "static" / "index.html").read_text(encoding="utf-8")
    assert 'id: "defs"' in js and 'id: "guide"' in js
    assert "defKeyFor" in js and "infoBtn" in js and "renderDefs" in js
    assert "/api/glossary" in js and "/api/selfcheck" in js and "/api/guide" in js
    assert 'id="opt-offline"' in html and 'id="env-panel"' in html


def test_glossary_lookup_maps_ui_labels():
    from src.metric_glossary import GLOSSARY, lookup

    assert lookup("MC Success") == "prob_terminal_gt_start"
    assert lookup("VaR 95% (monthly, historical)") == "var"
    assert lookup("Effective number of independent bets") == "effective_n"
    assert lookup("prob_all_payments_funded") == "prob_all_payments_funded"
    assert lookup("something not in the report at all") is None
    for key in list(GLOSSARY)[:10]:
        assert lookup(GLOSSARY[key]["label"]) == key


def test_tutorial_lists_every_tab_the_ui_actually_renders():
    """The tutorial's tab list and the dashboard cannot drift apart."""
    import re

    js = (PKG_ROOT / "web" / "static" / "app.js").read_text(encoding="utf-8")
    labels = re.findall(r'\{\s*id:\s*"[a-z_]+",\s*label:\s*"([^"]+)"', js)
    assert len(labels) >= 12, f"could not read the tab list from app.js (found {len(labels)})"

    text = " ".join((DOCS / "TUTORIAL.md").read_text(encoding="utf-8").split())
    missing = [label for label in labels if label not in text]
    assert not missing, f"TUTORIAL.md does not mention these tabs the UI renders: {missing}"


def test_tutorial_documents_the_interpretability_features():
    text = " ".join((DOCS / "TUTORIAL.md").read_text(encoding="utf-8").split())
    for must in ("What this means for your portfolio", "Chart captions", "Correlated clusters"):
        assert must in text, f"TUTORIAL.md does not document {must}"


def test_definitions_page_is_not_stale():
    """CI checks this too; catching it here means a stale page fails the suite, not just the build."""
    sys.path.insert(0, str(PKG_ROOT / "scripts"))
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "build_definitions_doc", PKG_ROOT / "scripts" / "build_definitions_doc.py")
    gen = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gen)

    expected = gen.build_doc()
    actual = (DOCS / "METRIC_DEFINITIONS.md").read_text(encoding="utf-8")
    assert actual == expected, (
        "docs/METRIC_DEFINITIONS.md is out of date - run "
        "`python portfolio_model/scripts/build_definitions_doc.py`")


def test_committed_sample_output_opens_and_is_documented():
    """A reader should be able to see what the deliverables look like before running anything."""
    from docx import Document
    from openpyxl import load_workbook

    sample_xlsx = DOCS / "sample_output" / "sample_output.xlsx"
    sample_docx = DOCS / "sample_output" / "sample_output_report.docx"
    assert sample_xlsx.exists(), "no sample workbook committed under docs/sample_output/"
    assert sample_docx.exists(), "no sample report committed under docs/sample_output/"

    wb = load_workbook(sample_xlsx)
    assert "Read Me" in wb.sheetnames and "Worked Examples" in wb.sheetnames
    assert len(wb.sheetnames) >= 20, f"sample workbook has only {len(wb.sheetnames)} sheets"

    doc = Document(sample_docx)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any(h.startswith("Appendix A") for h in headings), "sample report has no definitions appendix"
    assert any(h.startswith("Appendix B") for h in headings), "sample report has no worked examples"
    assert any(h.startswith("Appendix C") for h in headings), "sample report has no position guide"

    readme = (REPO_ROOT / "README.md").read_text(encoding="utf-8")
    assert "sample_output" in readme, "README does not point at the sample deliverables"


def test_tutorial_numbers_match_the_committed_sample_output():
    """The tutorial quotes real figures from the sample workbook. If the sample is regenerated
    and the tutorial is not updated, this fails rather than leaving a plausible-looking lie."""
    from openpyxl import load_workbook

    wb = load_workbook(DOCS / "sample_output" / "sample_output.xlsx", data_only=True)
    tutorial = " ".join((DOCS / "TUTORIAL.md").read_text(encoding="utf-8").split())

    def comparison_row(metric: str) -> list[float]:
        ws = wb["Model Comparison"]
        for r in ws.iter_rows(max_col=6, values_only=True):
            if r and str(r[0]).strip() == metric:
                return [float(v) for v in r[1:4] if isinstance(v, (int, float))]
        raise AssertionError(f"sample workbook has no {metric!r} row")

    money = lambda v: f"${v:,.0f}"
    percent = lambda v: f"{v * 100:.1f}%"

    med, p5, p95 = comparison_row("Median terminal value (2033)")[0], \
        comparison_row("5th percentile")[0], comparison_row("95th percentile")[0]
    funded = comparison_row("Probability reserve fully funded")

    expected = [
        money(med), money(p5), money(p95),
        percent(funded[0]), percent(funded[2]),
    ]

    ws = wb["Risk"]
    hhi = None
    for r in ws.iter_rows(max_col=4, values_only=True):
        if r and str(r[0]).strip() == "HHI raw":
            hhi = float(r[1])
    assert hhi is not None, "sample workbook has no HHI raw row"
    expected.append(f"{hhi:.4f}")

    for flag_id in ("POS-01", "RESV-03"):
        row = next((r for r in wb["Rule Flags"].iter_rows(max_col=8, values_only=True)
                    if r and str(r[0]).strip() == flag_id), None)
        assert row is not None, f"sample workbook has no {flag_id} flag"
        # the tutorial quotes each value rounded to 4 decimals, exactly as the report prints it
        expected += [f"{float(row[5]):.4f}", f"{float(row[6]):.4f}"]

    missing = [text for text in expected if text not in tutorial]
    assert not missing, (
        f"TUTORIAL.md quotes numbers that are not in the committed sample output: {missing}. "
        "Regenerate the tutorial section from the sample, or regenerate the sample.")
