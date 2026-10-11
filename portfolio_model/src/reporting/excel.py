"""Final Excel workbook (spec sections 81-96, 106)."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
KPI_FILL = PatternFill("solid", fgColor="DDEBF7")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, title: str | None = None,
              pct_cols: tuple[str, ...] = (), money_cols: tuple[str, ...] = ()) -> int:
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = TITLE_FONT
        r += 2
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=r, column=j, value=str(col))
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.border = THIN
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (_, row) in enumerate(df.iterrows(), start=r + 1):
        for j, col in enumerate(df.columns, start=1):
            v = row[col]
            if isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = None if np.isnan(v) else float(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            elif isinstance(v, pd.Timestamp):
                v = v.strftime("%Y-%m-%d")
            c = ws.cell(row=i, column=j, value=v)
            c.border = THIN
            if col in pct_cols and isinstance(v, (int, float)):
                c.number_format = "0.0%"
            elif col in money_cols and isinstance(v, (int, float)):
                c.number_format = "#,##0"
    return r + len(df) + 2


def _autosize(ws, max_w=42):
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_w, max(10, length + 2))


def build_workbook(path: Path, results: dict) -> Path:
    wb = Workbook()

    # ---- Sheet 1: Executive Summary / Dashboard (spec 81, 106)
    ws = wb.active
    ws.title = "Executive Summary"
    ws.cell(row=1, column=1, value="Portfolio Analysis Dashboard").font = Font(bold=True, size=16, color="1F4E79")
    ws.cell(row=2, column=1, value=f"Generated: {results['meta']['generated_utc']} | Seed: {results['meta']['seed']} | "
            f"Simulations: {results['meta']['n_sims']} | Model version: {results['meta']['version']}").font = Font(italic=True, size=9)
    kpis = results["dashboard"]
    r = 4
    for i, (label, value, fmt) in enumerate(kpis):
        row_i, col_i = r + (i // 4) * 3, 1 + (i % 4) * 2
        ws.cell(row=row_i, column=col_i, value=label).font = Font(bold=True, size=9, color="404040")
        c = ws.cell(row=row_i + 1, column=col_i, value=value)
        c.font = Font(bold=True, size=13)
        c.fill = KPI_FILL
        c.border = THIN
        if fmt:
            c.number_format = fmt

    # ---- Sheet 2: Holdings
    ws = wb.create_sheet("Holdings")
    _write_df(ws, results["holdings_sheet"], 1, "Holdings",
              pct_cols=("weight", "target_weight", "difference", "volatility", "expected_return",
                        "beta", "var_95", "cvar_95", "return_12m", "contribution_12m"),
              money_cols=("market_value",))
    _autosize(ws)

    # ---- Sheet 3: Sector Allocation
    ws = wb.create_sheet("Sector Allocation")
    _write_df(ws, results["sector_sheet"], 1, "Sector Allocation & Attribution",
              pct_cols=("weight", "target_weight", "difference", "sector_return_12m", "contribution_12m"))
    _autosize(ws)

    # ---- Sheet 4: Correlation Matrix (conditional formatting)
    ws = wb.create_sheet("Correlation Matrix")
    corr = results["corr"]
    ws.cell(row=1, column=1, value="Asset Correlation Matrix (historical monthly returns)").font = TITLE_FONT
    for j, t in enumerate(corr.columns, start=2):
        ws.cell(row=3, column=j, value=t).font = HDR_FONT
        ws.cell(row=3, column=j).fill = HDR_FILL
    for i, t in enumerate(corr.index, start=4):
        ws.cell(row=i, column=1, value=t).font = Font(bold=True)
        for j in range(len(corr.columns)):
            ws.cell(row=i, column=j + 2, value=float(corr.iloc[i - 4, j])).number_format = "0.00"
    n = len(corr)
    rng = f"B4:{get_column_letter(n + 1)}{n + 3}"
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num", start_value=-1, start_color="B71C1C",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=1, end_color="1A5276"))

    # ---- Sheet 5: Historical Performance
    ws = wb.create_sheet("Historical Performance")
    _write_df(ws, results["history_sheet"], 1, "Monthly Portfolio History",
              pct_cols=("portfolio_return", "cumulative_return", "drawdown", "rolling_vol_12m", "rolling_sharpe_12m"))
    _autosize(ws)

    # ---- Sheet 6: Monte Carlo
    ws = wb.create_sheet("Monte Carlo")
    ws.cell(row=1, column=1, value=f"Monte Carlo — {results['mc']['model']} model, {results['mc']['n_sims']:,} simulations, seed {results['meta']['seed']}").font = TITLE_FONT
    _write_df(ws, results["mc_sheet"], 3, "Simulation Summary (sample of 200 trials)",
              money_cols=("terminal_value", "min_value"), pct_cols=("cagr", "max_drawdown", "success"))
    _write_df(ws, results["mc_percentiles"], 3 + len(results["mc_sheet"]) + 3, "Percentile Summary",
              money_cols=("value",))
    _autosize(ws)

    # ---- Sheet 7: Markov Model
    ws = wb.create_sheet("Markov Model")
    ws.cell(row=1, column=1, value="Markov Regime Model").font = TITLE_FONT
    r = _write_df(ws, results["markov_states"], 3, "State Definitions & Statistics",
                  pct_cols=("mean_return", "std_return"))
    r = _write_df(ws, results["markov_P"], r + 1, "Transition Matrix P(i->j)",
                  pct_cols=tuple(results["markov_P"].columns))
    r = _write_df(ws, results["markov_misc"], r + 1, "Stationary Distribution / Durations / Forecasts")
    _autosize(ws)

    # ---- Sheet 8: Historical Bootstrap
    ws = wb.create_sheet("Historical Bootstrap")
    ws.cell(row=1, column=1, value=f"Historical Bootstrap — {results['bs']['method']} blocks of {results['bs']['block_size_days']} trading days, {results['bs']['n_sims']:,} simulations").font = TITLE_FONT
    _write_df(ws, results["bs_sheet"], 3, "Simulation Summary (sample of 200 trials)",
              money_cols=("terminal_value",), pct_cols=("cagr", "max_drawdown"))
    _autosize(ws)

    # ---- Sheet 9: Model Comparison
    ws = wb.create_sheet("Model Comparison")
    end = _write_df(ws, results["comparison_sheet"], 1, "Monte Carlo vs Markov vs Historical Bootstrap",
                    money_cols=("Monte Carlo", "Markov", "Historical Bootstrap"))
    _write_df(ws, results["disagreement_sheet"], end, "Model Disagreement Analysis")
    _autosize(ws)

    # ---- Sheet 10: Sensitivity
    ws = wb.create_sheet("Sensitivity")
    _write_df(ws, results["sensitivity_sheet"], 1, "Return x Volatility Sensitivity Grid",
              pct_cols=("return_shift", "prob_success", "prob_all_payments_funded"),
              money_cols=("median_terminal", "p5_terminal"))
    _autosize(ws)

    # ---- Sheet 11: Risk
    ws = wb.create_sheet("Risk")
    end = _write_df(ws, results["risk_sheet"], 1, "Portfolio Risk Metrics",
                    pct_cols=("value",))
    end = _write_df(ws, results["recovery_sheet"], end, "Major Drawdown Recovery Events",
                    pct_cols=("drawdown",))
    clusters = results.get("cluster_sheet")
    if clusters is not None and len(clusters):
        _write_df(ws, clusters, end, "Correlated Clusters (holdings that are really one bet)",
                  pct_cols=("combined_weight", "mean_pairwise_correlation"))
    _autosize(ws)

    # ---- Sheet 12: Stress Tests
    ws = wb.create_sheet("Stress Tests")
    end = _write_df(ws, results["stress_sheet"], 1, "Historical Crisis Stress Tests",
                    pct_cols=("portfolio_return", "max_drawdown", "worst_day", "worst_week", "worst_month"))
    _write_df(ws, results["hypothetical_sheet"], end,
              "Hypothetical Forward Scenarios (assumptions, not forecasts)", pct_cols=("portfolio_return",))
    _autosize(ws)

    # ---- Sheet 13: Glide Path
    ws = wb.create_sheet("Glide Path")
    end = _write_df(ws, results["glide_sheet"], 1, "2033 Glide Path",
                    pct_cols=("equity", "bond", "cash"))
    _write_df(ws, results["glide_impact_sheet"], end,
              "Glide Path Impact Comparison", money_cols=("median_terminal", "p5_terminal"),
              pct_cols=("prob_depletion", "prob_payments_funded"))
    _autosize(ws)

    # ---- Sheet 14: Rebalancing
    ws = wb.create_sheet("Rebalancing")
    _write_df(ws, results["rebalance_sheet"], 1, "Rebalancing Bands & Suggested Trades",
              pct_cols=("current_weight", "target", "lower_band", "upper_band", "suggested_trade", "turnover"))
    _autosize(ws)

    # ---- Sheet 14b: Rule Flags (deterministic, threshold-based)
    ws = wb.create_sheet("Rule Flags")
    rule_df = pd.DataFrame(results.get("rule_flags", []))
    if len(rule_df):
        rule_df = rule_df[["id", "severity", "category", "metric", "current_value", "threshold",
                           "rule", "action", "evidence"]]
        end = _write_df(ws, rule_df, 1, "Rule-Based Risk & Rebalancing Flags (deterministic; thresholds from config)")
        ws.cell(row=end, column=1,
                value="These are rule-based portfolio risk flags, not AI-generated recommendations.").font = Font(italic=True, size=9)
    else:
        ws.cell(row=1, column=1, value="No rule flags triggered.").font = TITLE_FONT
    _autosize(ws)

    # ---- Sheet 15: Investment Theses
    ws = wb.create_sheet("Investment Theses")
    _write_df(ws, results["thesis_sheet"], 1, "Investment Theses (spec section 53)")
    _autosize(ws)

    # ---- Sheet 16: Data Quality
    ws = wb.create_sheet("Data Quality")
    _write_df(ws, results["dq_sheet"], 1, "Data Quality Flags & Limitations")
    for row in ws.iter_rows(min_row=4, max_col=4):
        for c in row:
            c.fill = WARN_FILL
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws)

    # ---- Sheet 17: Scorecard & Decision Matrix
    ws = wb.create_sheet("Scorecard")
    end = _write_df(ws, results["scorecard"], 1, "Portfolio Scorecard (documented heuristic scoring)")
    _write_df(ws, results["decision_matrix"], end, "Final Decision Matrix")
    _autosize(ws)

    # ---- Sheet 18: Stop-Loss Backtests
    ws = wb.create_sheet("Stop-Loss Backtests")
    _write_df(ws, results["stop_sheet"], 1, "Stop-Loss Backtests (historical, not advice)",
              pct_cols=("cagr", "vol", "max_drawdown", "total_return", "missed_upside"))
    _autosize(ws)

    # ---- Sheet 19: Metric Definitions (rendered from src/metric_glossary.py)
    from ..metric_glossary import GLOSSARY

    ws = wb.create_sheet("Definitions")
    def_rows = [{"key": k, "metric": e["label"], "unit": e["unit"], "what_it_is": e["formula"],
                 "what_it_measures": e["meaning"], "how_to_read_it": e["read_as"],
                 "what_to_do": e["action"], "trap": e["caveat"]}
                for k, e in GLOSSARY.items()]
    end = _write_df(ws, pd.DataFrame(def_rows), 1,
                    "Metric Definitions — what every number in this workbook means")
    ws.cell(row=end, column=1,
            value="Narrative version with the ten most common misreadings: docs/METRIC_DEFINITIONS.md "
                  "(generated from the same source file).").font = Font(italic=True, size=9)
    for row in ws.iter_rows(min_row=4, max_col=8):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws, max_w=58)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
