"""Spreadsheet ingestion with canonical-schema detection (spec sections 3-4)."""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl
import pandas as pd

# canonical field <- possible input names (lower-cased, normalized)
ALIASES: dict[str, list[str]] = {
    "ticker": ["ticker", "symbol", "security", "company", "name", "name / description"],
    "quantity": ["shares", "quantity", "units", "share count"],
    "price": ["current price", "price", "last price", "market price"],
    "cost_basis": ["cost basis", "avg cost", "average cost", "cost"],
    "market_value": ["market value", "current value", "value", "position value"],
    "portfolio_weight": ["portfolio weight", "weight", "% of portfolio", "pre-cap weight (k)"],
    "target_weight": ["target weight", "target allocation"],
    "bucket_target": ["bucket target (h)", "bucket target"],
    "sector": ["sector", "gics sector", "industry sector"],
    "industry": ["industry", "sub industry", "sub-industry"],
    "asset_class": ["asset class", "class", "bucket", "type"],
    "country": ["country", "region"],
    "currency": ["currency", "ccy"],
    "purchase_date": ["purchase date", "buy date", "date"],
    "stop_loss": ["stop loss", "stop-loss", "stop"],
    "notes": ["notes", "comment", "commentary"],
    "thesis": ["investment thesis", "thesis", "thematic rationale", "why own"],
    "dividend_yield": ["dividend yield", "yield"],
    "expected_return": ["expected return", "exp return", "mean return"],
    "volatility": ["volatility (std dev)", "volatility", "std dev", "stdev", "risk"],
    "correlation_group": ["correlation group", "correlation group (blank = independent)", "group", "cluster"],
}

TICKER_RE = re.compile(r"^[A-Z][A-Z.\-]{0,9}$")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def _map_header(header: list) -> dict[int, str]:
    mapping = {}
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        n = _norm(cell)
        for canon, aliases in ALIASES.items():
            if n in aliases or any(a in n for a in aliases if len(a) > 4):
                if canon not in mapping.values():
                    mapping[idx] = canon
                    break
    return mapping


@dataclass
class IngestResult:
    holdings: pd.DataFrame
    detected_columns: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    source_file: str = ""


def _find_holdings_table(ws) -> tuple[int, int, dict[int, str]] | None:
    """Scan a worksheet for a header row that contains a ticker-like column."""
    max_scan = min(ws.max_row, 60)
    for r in range(1, max_scan + 1):
        row = [c.value for c in ws[r]]
        mapping = _map_header(row)
        canon = set(mapping.values())
        if "ticker" in canon and ({"quantity", "market_value", "portfolio_weight", "target_weight"} & canon):
            return r, ws.max_column, mapping
    return None


def ingest(path: str, sheet: str | None = None) -> IngestResult:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = [sheet] if sheet else wb.sheetnames
    frames: list[pd.DataFrame] = []
    detected: dict[str, str] = {}
    warnings: list[str] = []

    for name in sheets:
        ws = wb[name]
        found = _find_holdings_table(ws)
        if not found:
            continue
        header_row, max_col, mapping = found
        rows = []
        for r in range(header_row + 1, ws.max_row + 1):
            vals = [ws.cell(row=r, column=c + 1).value for c in range(max_col)]
            rec = {canon: vals[i] for i, canon in mapping.items()}
            t = rec.get("ticker")
            if t is None:
                continue
            t = str(t).strip()
            # skip subtotal/label rows
            if not TICKER_RE.match(t.upper().replace(" ", "")):
                continue
            if t.upper() in {"TOTAL", "SUBTOTAL", "SUM", "CASH TOTAL", "TBD"}:
                continue
            if any(k in _norm(t) for k in ("subtotal", "total", "your", "pick", "check")):
                continue
            rec["ticker"] = t.upper()
            if all(v is None for k, v in rec.items() if k != "ticker"):
                continue
            rows.append(rec)
        if rows:
            df = pd.DataFrame(rows)
            df["_source_sheet"] = name
            frames.append(df)
            for canon in df.columns:
                detected.setdefault(canon, name)

    if not frames:
        raise ValueError(f"No holdings table detected in {path}. Expected a header row with a ticker column plus quantity/value/weight.")

    holdings = pd.concat(frames, ignore_index=True)
    # merge duplicate detections across sheets (prefer non-null)
    holdings = holdings.groupby("ticker", as_index=False).first()

    warnings.append(f"Detected columns: {sorted(k for k in holdings.columns if k != '_source_sheet')}")
    return IngestResult(holdings=holdings, detected_columns=detected, warnings=warnings, source_file=str(path))
